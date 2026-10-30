"""PDF and Excel builders (no Django imports, so they are easy to test).

Used for: result slips (RES-06), credential sheets (AUTH-06), subject reports (TCH-09)
and admin exports (ADM-11)."""
import io
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BRAND = colors.HexColor("#1F6F5C")
LIGHT = colors.HexColor("#E8F3F0")
_styles = getSampleStyleSheet()
H1 = ParagraphStyle("h1", parent=_styles["Title"], textColor=BRAND, fontSize=18, spaceAfter=4)
H2 = ParagraphStyle("h2", parent=_styles["Heading2"], textColor=BRAND, fontSize=12, spaceBefore=10, spaceAfter=4)
BODY = ParagraphStyle("body", parent=_styles["BodyText"], fontSize=9.5, leading=12)
SMALL = ParagraphStyle("small", parent=BODY, fontSize=8, textColor=colors.grey)
CELL = ParagraphStyle("cell", parent=BODY, fontSize=8.5, leading=10.5)
CELL_HEAD = ParagraphStyle("cellh", parent=CELL, textColor=colors.white, fontName="Helvetica-Bold")


_LOGO = Path(__file__).resolve().parent / "assets" / "logo.png"


def _p(text, style=CELL):
    return Paragraph(escape("" if text is None else str(text)), style)


def _grid(headers, rows, col_widths=None):
    data = [[_p(h, CELL_HEAD) for h in headers]] + [[_p(c) for c in r] for r in rows]
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BRAND),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B7C9C4")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return t


def _build(story, title, footer, wide=False):
    buf = io.BytesIO()
    size = landscape(A4) if wide else A4
    doc = SimpleDocTemplate(buf, pagesize=size, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=22 * mm, bottomMargin=18 * mm, title=title)

    def decorate(canvas, d):
        canvas.saveState()
        if _LOGO.exists():  # school crest, top right of every page
            canvas.drawImage(str(_LOGO), size[0] - 15 * mm - 16 * mm, size[1] - 5 * mm - 16 * mm, 16 * mm, 16 * mm, mask="auto")
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(15 * mm, 10 * mm, footer)
        canvas.drawRightString(size[0] - 15 * mm, 10 * mm, f"Page {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buf.getvalue()


def sections_pdf(school, title, subtitle, sections, wide=True):
    """sections: list of {"heading": str, "headers": [...], "rows": [[...]]}"""
    story = [Paragraph(escape(school), H1), Paragraph(escape(title), H2)]
    if subtitle:
        story.append(Paragraph(escape(subtitle), SMALL))
    for sec in sections:
        story.append(Paragraph(escape(sec["heading"]), H2))
        if sec["rows"]:
            story.append(_grid(sec["headers"], sec["rows"]))
        else:
            story.append(Paragraph("No data yet.", SMALL))
        story.append(Spacer(1, 4 * mm))
    return _build(story, title, f"{school} - {title}", wide=wide)


def credentials_pdf(school, groups):
    """groups: list of (class_label, [(name, username, password), ...]). One page group per class."""
    story = []
    for i, (label, rows) in enumerate(groups):
        if i:
            story.append(PageBreak())
        story += [
            Paragraph(escape(school), H1),
            Paragraph(f"Login details - {escape(label)}", H2),
            Paragraph("Keep this sheet private. Log in with your username and the starting password in the last column, "
                      "then choose a new password of your own. The starting password stops working if it is not used in time.", SMALL),
            Spacer(1, 3 * mm),
            _grid(["#", "Name", "Username", "Starting password"],
                  [[n, name, user, pw] for n, (name, user, pw) in enumerate(rows, 1)],
                  col_widths=[10 * mm, 70 * mm, 45 * mm, 45 * mm]),
        ]
    if not groups:
        story.append(Paragraph("No students to list.", BODY))
    return _build(story, "Login details", f"{school} - login details (confidential)")


def result_slip_pdf(school, d):
    """d keys: student, klass, exam, subject, date, score, total, percent, grade, points, meaning,
    passed, time_taken, class_average, position (or None), topics [(topic, percent)], revise [str]."""
    mins, secs = divmod(int(d.get("time_taken") or 0), 60)
    info = Table([
        [_p("Student", CELL_HEAD), _p(d["student"]), _p("Class", CELL_HEAD), _p(d["klass"])],
        [_p("Exam", CELL_HEAD), _p(d["exam"]), _p("Subject", CELL_HEAD), _p(d["subject"])],
        [_p("Date", CELL_HEAD), _p(d["date"]), _p("Time taken", CELL_HEAD), _p(f"{mins} min {secs} s")],
    ], colWidths=[28 * mm, 62 * mm, 28 * mm, 62 * mm])
    info.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), BRAND), ("BACKGROUND", (2, 0), (2, -1), BRAND),
                              ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B7C9C4"))]))
    big = ParagraphStyle("big", parent=BODY, fontSize=22, leading=26, textColor=BRAND, alignment=1)
    result = Table([[
        Paragraph(f"{d['score']:g} / {d['total']:g}", big),
        Paragraph(f"{d['percent']:g}%", big),
        Paragraph(f"Grade {escape(str(d['grade']))}", big),
    ], [_p("Score"), _p("Percentage"), _p(f"{d.get('meaning', '')} - {d['points']} point(s)")]],
        colWidths=[60 * mm, 60 * mm, 60 * mm])
    result.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("BOX", (0, 0), (-1, -1), 1, BRAND),
                                ("BACKGROUND", (0, 0), (-1, -1), LIGHT)]))
    story = [Paragraph(escape(school), H1), Paragraph("Result slip", H2), info, Spacer(1, 5 * mm), result,
             Spacer(1, 3 * mm)]
    extra = [f"Result: {'Pass' if d['passed'] else 'Below the pass mark'}",
             f"Class average (anonymous): {d['class_average']:g}%"]
    if d.get("position"):
        extra.append(f"Class position: {d['position']}")
    story.append(Paragraph(escape(" | ".join(extra)), BODY))
    if d.get("topics"):
        story += [Paragraph("Performance by topic", H2),
                  _grid(["Topic", "Score"], [[t, f"{p:g}%"] for t, p in d["topics"]],
                        col_widths=[120 * mm, 40 * mm])]
    if d.get("revise"):
        story += [Paragraph("What to revise", H2), Paragraph(escape(", ".join(d["revise"])), BODY)]
    story += [Spacer(1, 8 * mm),
              Paragraph("This slip shows your own private result. Grades follow the school grading scale in force "
                        "when the exam was marked.", SMALL)]
    return _build(story, "Result slip", f"{school} - result slip")


# ------------------------------------------------------------------ Excel

def _sheet_title(t, used):
    t = "".join(c for c in str(t) if c not in '[]:*?/\\')[:31] or "Sheet"
    base, n = t, 2
    while t in used:
        t = f"{base[:28]} {n}"
        n += 1
    used.add(t)
    return t


def sheets_xlsx(sheets):
    """sheets: list of (title, headers, rows). Text is stored as text, so a name that starts
    with '=' can never run as a formula."""
    wb, used = Workbook(), set()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(_sheet_title(title, used))
        ws.append(list(headers))
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F6F5C")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(list(r))
            for c in ws[ws.max_row]:
                if isinstance(c.value, str):
                    c.data_type = "s"
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r)] or [8])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 50)
    if not wb.sheetnames:
        wb.create_sheet("Report")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_safe(v):
    """Stop spreadsheet formula injection in CSV downloads."""
    s = "" if v is None else str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@") else s
