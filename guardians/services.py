"""Parents: file upload, linking, and queuing the SMS that carry each student's login details."""
import csv
import io
import json
import re
import secrets
from datetime import timedelta

from django.core.exceptions import ImproperlyConfigured
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.authtoken.models import Token

from accounts.models import User
from . import crypto
from .models import Guardian, GuardianLink, SmsConfig, SmsJob
from .phones import mask, normalize_phone

TEMPLATE_COLUMNS = ["parent_name", "phone", "relationship", "student_index", "whatsapp", "language"]
ALIASES = {
    "parent_name": "parent_name", "name": "parent_name", "parent": "parent_name", "guardian": "parent_name", "guardian_name": "parent_name",
    "phone": "phone", "phone_number": "phone", "mobile": "phone", "simu": "phone", "telephone": "phone",
    "relationship": "relationship", "relation": "relationship",
    "student_index": "student_index", "index": "student_index", "index_number": "student_index", "student": "student_index", "students": "student_index",
    "whatsapp": "whatsapp", "has_whatsapp": "whatsapp", "language": "language", "lang": "language",
}
MAX_ROWS = 3000
PASSWORD_ALPHABET = "abcdefghjkmnpqrstuvwxyz23456789"  # no 0/o, 1/l/i, so it can be read from an SMS
MAX_ATTEMPTS = 3
LEASE_MINUTES = 10
NOTE_MAX = 160  # the optional custom line an administrator may add to every message
_GSM_BASIC = set("@£$¥èéùìòÇ\nØø\rÅåΔ_ΦΓΛΩΠΨΣΘΞÆæßÉ !\"#¤%&'()*+,-./0123456789:;<=>?¡ABCDEFGHIJKLMNOPQRSTUVWXYZÄÖÑÜ§¿abcdefghijklmnopqrstuvwxyzäöñüà")
_GSM_EXT = set("^{}\\[~]|€")
_PUNCT = {"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u00a0": " "}


# ------------------------------------------------------------------ reading the file
def parse_file(f):
    """Return a list of dicts using the template column names. Accepts .csv and .xlsx."""
    name = (getattr(f, "name", "") or "").lower()
    if name.endswith(".xlsx"):
        ws = load_workbook(f, read_only=True, data_only=True).active
        table = [list(r) for r in ws.iter_rows(values_only=True)]
    elif name.endswith(".csv") or name.endswith(".txt"):
        text = f.read().decode("utf-8-sig", errors="replace")
        table = list(csv.reader(io.StringIO(text)))
    else:
        raise ValueError("Please upload the template as .csv or .xlsx.")
    table = [r for r in table if any(str(c or "").strip() for c in r)]
    if not table:
        raise ValueError("The file is empty.")
    head = [ALIASES.get(re.sub(r"[\s\-]+", "_", str(h or "").strip().lower())) for h in table[0]]
    if "phone" not in head or "student_index" not in head:
        raise ValueError("The first row must have the columns: parent_name, phone, relationship, student_index, whatsapp, language.")
    if len(table) - 1 > MAX_ROWS:
        raise ValueError(f"Please upload at most {MAX_ROWS} rows at a time.")
    rows = []
    for n, r in enumerate(table[1:], start=2):
        d = {"_row": n}
        for key, val in zip(head, r):
            if key and key not in d:
                d[key] = val
        rows.append(d)
    return rows


def _yes(v):
    return str(v or "").strip().lower() in ("yes", "y", "true", "1", "ndio", "x")


def _clean_name(v):
    return re.sub(r"\s+", " ", re.sub(r"[\x00-\x1f<>]", " ", str(v or ""))).strip()[:120]


def split_indexes(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return [x for x in re.split(r"[,;/\s]+", str(v or "").strip()) if x]


def student_lookup():
    qs = User.objects.filter(role=User.Role.STUDENT, is_active=True, is_archived=False).exclude(index_number="").select_related("school_class")
    out = {}
    for u in qs:
        out.setdefault(u.index_number.strip().lower(), u)
    return out


def validate_rows(rows):
    """Check every row. Nothing is saved. Returns (clean_rows, report)."""
    students = student_lookup()
    clean, report_rows, seen_phone = [], [], {}
    for d in rows:
        problems, warns = [], []
        name = _clean_name(d.get("parent_name"))
        phone = normalize_phone(d.get("phone"))
        if not name:
            problems.append("Parent name is missing.")
        if not phone:
            problems.append("Phone number is not a valid Tanzanian mobile number.")
        found, unknown = [], []
        idx = split_indexes(d.get("student_index"))
        if not idx:
            problems.append("Student index number is missing.")
        for i in idx:
            u = students.get(i.lower())
            (found if u else unknown).append(u or i)
        if unknown:
            problems.append("No active student with index number: " + ", ".join(unknown) + ".")
        lang = str(d.get("language") or "").strip().lower()[:2]
        if lang and lang not in ("sw", "en"):
            warns.append("Language should be sw or en; Kiswahili will be used.")
        if phone and phone in seen_phone:
            warns.append(f"Same phone as row {seen_phone[phone]}; the children are joined to one parent.")
        elif phone:
            seen_phone[phone] = d["_row"]
        status = "error" if problems else ("warn" if warns else "ok")
        report_rows.append({"row": d["_row"], "name": name, "phone": mask(phone) if phone else str(d.get("phone") or ""),
                            "students": [f"{u.first_name} {u.last_name} ({u.index_number})" for u in found],
                            "status": status, "problems": problems + warns})
        if not problems:
            clean.append({"name": name, "phone": phone, "relationship": _clean_name(d.get("relationship"))[:30],
                          "student_ids": [u.pk for u in found], "whatsapp": _yes(d.get("whatsapp")),
                          "language": lang if lang in ("sw", "en") else "sw"})
    phones = {c["phone"] for c in clean}
    existing = set(Guardian.objects.filter(phone_idx__in=[crypto.blind_index(p) for p in phones]).values_list("phone_idx", flat=True))
    new_links = 0
    for c in clean:
        g_exists = crypto.blind_index(c["phone"]) in existing
        gid = Guardian.objects.filter(phone_idx=crypto.blind_index(c["phone"])).values_list("pk", flat=True).first() if g_exists else None
        for sid in c["student_ids"]:
            if not (gid and GuardianLink.objects.filter(guardian_id=gid, student_id=sid).exists()):
                new_links += 1
    report = {"rows": report_rows, "total_rows": len(rows), "valid_rows": len(clean), "invalid_rows": len(rows) - len(clean),
              "parents_new": len(phones) - len(existing), "parents_existing": len(existing), "links_new": new_links}
    return clean, report


@transaction.atomic
def commit_rows(clean, user):
    created = updated = links = 0
    for c in clean:
        g = Guardian.objects.filter(phone_idx=crypto.blind_index(c["phone"])).first()
        if g is None:
            g = Guardian(created_by=user, relationship=c["relationship"], has_whatsapp=c["whatsapp"], language=c["language"])
            g.set_name(c["name"])
            g.set_phone(c["phone"])
            g.save()
            created += 1
        else:
            changed = False
            if c["whatsapp"] and not g.has_whatsapp:
                g.has_whatsapp, changed = True, True
            if c["relationship"] and not g.relationship:
                g.relationship, changed = c["relationship"], True
            if changed:
                g.save()
                updated += 1
        for sid in c["student_ids"]:
            _, made = GuardianLink.objects.get_or_create(guardian=g, student_id=sid)
            links += made
    return {"parents_created": created, "parents_updated": updated, "links_created": links}


# ------------------------------------------------------------------ parents given in the STUDENT template
# The student template (accounts app) may carry up to two parents per student in optional columns. Everything below is only used
# by that upload; the parents' own template above is unchanged. Parents are found by phone number, so a parent with several
# children (in one file, in several files, or already registered) is always ONE Guardian linked to all of them.
PARENT_SLOTS = (
    {"label": "Parent",
     "name": ("parent name", "parent", "guardian name", "guardian", "parent's name"),
     "phone": ("parent phone", "parent phone number", "parent mobile", "parent contact", "parent telephone", "guardian phone"),
     "relationship": ("parent relationship", "relationship", "relation", "guardian relationship")},
    {"label": "Parent 2",
     "name": ("parent 2 name", "parent2 name", "second parent name", "parent 2"),
     "phone": ("parent 2 phone", "parent2 phone", "parent 2 phone number", "parent 2 mobile", "second parent phone"),
     "relationship": ("parent 2 relationship", "parent2 relationship", "second parent relationship")},
)


def _cell(d, keys):
    for k in keys:
        v = str(d.get(k) or "").strip()
        if v:
            return v
    return ""


def _phone_cell(v):
    """normalize_phone, plus what Excel does to numbers typed into a cell (712345678.0 or 7.12345678E+8)."""
    s = str(v or "").strip()
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    elif re.fullmatch(r"\d(\.\d+)?[eE]\+?\d+", s):
        try:
            s = str(int(float(s)))
        except ValueError:
            pass
    return normalize_phone(s)


def _same_name(a, b):
    return re.sub(r"\s+", " ", a or "").strip().lower() == re.sub(r"\s+", " ", b or "").strip().lower()


def read_parent_cells(d):
    """The parents on one row of the student template. Returns (parents, problems, warnings).
    A parent is {"name", "phone", "relationship"}; the name may still be empty (another row or a registered parent can supply it)."""
    found, problems, warns, seen = [], [], [], set()
    for slot in PARENT_SLOTS:
        name = _clean_name(_cell(d, slot["name"]))
        raw = _cell(d, slot["phone"])
        rel = _clean_name(_cell(d, slot["relationship"]))[:30]
        if not (name or raw):
            continue  # nothing for this parent (a relationship on its own is ignored)
        if not raw:
            problems.append(f"{slot['label']}: the phone number is missing.")
            continue
        phone = _phone_cell(raw)
        if not phone:
            problems.append(f"{slot['label']}: the phone number is not a valid Tanzanian mobile number.")
            continue
        if phone in seen:
            warns.append(f"{slot['label']} has the same phone number as the other parent on this row; counted once.")
            continue
        seen.add(phone)
        found.append({"name": name, "phone": phone, "relationship": rel})
    return found, problems, warns


def ensure_parent_key():
    """Parents are stored encrypted, which needs GUARDIAN_KEY. Say so clearly now instead of failing at the confirm step."""
    try:
        crypto.blind_index("+255700000000")
    except ImproperlyConfigured:
        raise ValueError("Parents' phone numbers cannot be saved yet because GUARDIAN_KEY is not set on the server. "
                         "Set it first, or upload the students without the parent columns.")


def settle_parents(entries):
    """entries = [(row, [parent, ...])] for rows that are otherwise valid. Works out the final name of every phone number
    (a registered parent's name wins, then the first name written in the file), fills it into each parent, and returns
    (errors, warnings, stats) where errors/warnings are {row: [text]}. Nothing is saved."""
    errors, warns = {}, {}
    idx = {p["phone"]: crypto.blind_index(p["phone"]) for _, ps in entries for p in ps}
    registered = {g.phone_idx: g for g in Guardian.objects.filter(phone_idx__in=set(idx.values()))}
    first_name, first_rel = {}, {}
    for row, ps in entries:
        for p in ps:
            if p["name"] and p["phone"] not in first_name:
                first_name[p["phone"]] = (p["name"], row)
            if p["relationship"] and p["phone"] not in first_rel:
                first_rel[p["phone"]] = p["relationship"]
    names = {}
    for phone, key in idx.items():
        g = registered.get(key)
        names[phone] = (g.name, "a registered parent") if g and g.name else first_name.get(phone, ("", ""))
        if isinstance(names[phone][1], int):
            names[phone] = (names[phone][0], f"row {names[phone][1]}")
    for row, ps in entries:
        for p in ps:
            final, source = names[p["phone"]]
            if not final:
                errors.setdefault(row, []).append(f"Parent {mask(p['phone'])}: the name is missing.")
                continue
            if p["name"] and not _same_name(p["name"], final):
                warns.setdefault(row, []).append(
                    f"{mask(p['phone'])} already belongs to \"{final}\" ({source}); the student is linked to that parent. Check the number.")
            p["name"] = final
            p["relationship"] = p["relationship"] or first_rel.get(p["phone"], "")
    ok = [(r, ps) for r, ps in entries if r not in errors]
    phones = {p["phone"] for _, ps in ok for p in ps}
    known = sum(1 for ph in phones if idx[ph] in registered)
    per_phone = {}
    for _, ps in ok:
        for p in ps:
            per_phone[p["phone"]] = per_phone.get(p["phone"], 0) + 1
    stats = {"rows_with_parents": len(ok), "parents_new": len(phones) - known, "parents_existing": known,
             "links": sum(per_phone.values()), "parents_with_several_students": sum(1 for n in per_phone.values() if n > 1)}
    return errors, warns, stats


def encode_parents(parents):
    """Parents waiting for the admin's confirmation are kept encrypted, like everywhere else."""
    return crypto.encrypt(json.dumps(parents)) if parents else ""


def decode_parents(token):
    try:
        return json.loads(crypto.decrypt(token) or "[]") if token else []
    except ValueError:
        return []


@transaction.atomic
def link_new_students(pairs, user):
    """pairs = [(student id, [parent, ...])]. One Guardian per phone number (an existing one is reused), linked to every
    student that names it. Runs inside the student creation, so a failure here creates nothing."""
    by_phone = {}
    for sid, ps in pairs:
        for p in ps:
            c = by_phone.setdefault(p["phone"], {"name": "", "phone": p["phone"], "relationship": "", "student_ids": [],
                                                 "whatsapp": False, "language": "sw"})
            c["name"] = c["name"] or p["name"]
            c["relationship"] = c["relationship"] or p.get("relationship", "")
            if sid not in c["student_ids"]:
                c["student_ids"].append(sid)
    for c in by_phone.values():
        c["name"] = c["name"] or "Parent"  # only if the name vanished between checking and confirming
    return commit_rows(list(by_phone.values()), user) if by_phone else {"parents_created": 0, "parents_updated": 0, "links_created": 0}


# ------------------------------------------------------------------ the optional custom line
def clean_note(v):
    """The administrator's optional extra line: one line of plain text. Curly quotes and dashes (what phones and Word type) become
    plain ones so the SMS stays cheap, control characters and line breaks become spaces, extra spaces are squeezed."""
    v = "".join(_PUNCT.get(c, c) for c in str(v or ""))
    return re.sub(r"\s+", " ", re.sub(r"[\x00-\x1f\x7f<>]", " ", v)).strip()


def sms_parts(text):
    """(number of SMS the text will use, True if it has characters outside the cheap GSM alphabet, e.g. emoji)."""
    if all(c in _GSM_BASIC or c in _GSM_EXT for c in text):
        n = sum(2 if c in _GSM_EXT else 1 for c in text)
        return (1 if n <= 160 else -(-n // 153)), False
    n = len(text.encode("utf-16-le")) // 2
    return (1 if n <= 70 else -(-n // 67)), True


# ------------------------------------------------------------------ the message
def new_password():
    return "".join(secrets.choice(PASSWORD_ALPHABET) for _ in range(8))


DEFAULT_MESSAGES = {  # what the parent receives unless the administrator edits it. {url} = the school web address, {children} = one line per child
    "sw": "Amani Girls: taarifa za kuingia kwenye mitihani. Fungua {url}\n{children}\nWeka nenosiri jipya ukiingia mara ya kwanza. Usimpe mtu.",
    "en": "Amani Girls: exam login. Open {url}\n{children}\nSet a new password at first login. Do not share.",
}
CHILD_LINE = {"sw": "{n}: jina {u}, nenosiri {p}", "en": "{n}: user {u}, password {p}"}
TEMPLATE_MAX = 300  # the wording around the children's lines
CUSTOM_MAX = 459  # a message with no login details: up to 3 ordinary SMS


def clean_template(v):
    """Edited wording: keeps line breaks, turns curly quotes/dashes into plain ones, drops control characters."""
    v = "".join(_PUNCT.get(c, c) for c in str(v or "").replace("\r\n", "\n").replace("\r", "\n"))
    lines = [re.sub(r"[ \t]+", " ", re.sub(r"[\x00-\x09\x0b-\x1f\x7f<>]", " ", ln)).strip() for ln in v.split("\n")]
    return "\n".join(lines).strip("\n")


def check_templates(d):
    """{lang: edited text} from the request. A text that is empty or equals the standard wording is returned as "" (= use the standard)."""
    out = {}
    for lang in ("sw", "en"):
        text = clean_template(d.get("message_" + lang))
        if not text or text == clean_template(DEFAULT_MESSAGES[lang]):
            out[lang] = ""
            continue
        if len(text.replace("{children}", "").replace("{url}", "")) > TEMPLATE_MAX:
            raise ValueError(f"The message wording is too long. Keep the text around the login lines to {TEMPLATE_MAX} characters or fewer.")
        if "{children}" not in text:
            raise ValueError("The message must contain {children}: that is where each child's name, user name and password go. Without it parents would get no login details.")
        out[lang] = text
    return out


def build_body(lang, cfg, pairs, note="", template=""):
    """pairs = [(student display name, username, password)]; note = the optional custom line; template = edited wording (or "" for the
    standard one). Plain characters only, to keep it cheap (GSM-7)."""
    url = (cfg.public_url or "").strip() or "the school website"
    lang = lang if lang in DEFAULT_MESSAGES else "sw"
    kids = "\n".join(CHILD_LINE[lang].format(n=n, u=u, p=p) for n, u, p in pairs)
    text = (template or DEFAULT_MESSAGES[lang]).replace("{url}", url).replace("{children}", kids)
    return text + ("\n" + note if note else "")  # the administrator's own line goes last, after the security reminder


def preview(items, cfg, note="", templates=None, login=True):
    """What a message will look like (first parent in the list, with a placeholder instead of a real password), and how many SMS it uses."""
    if not items:
        return {"text": "", "parts": 0, "unicode": False}
    g, kids = items[0]
    if login:
        text = build_body(g.language, cfg, [(_display(s), s.username, "xxxxxxxx") for s in kids], note, (templates or {}).get(g.language, ""))
    else:
        text = note  # a custom message: exactly what the administrator wrote
    parts, uni = sms_parts(text)
    return {"text": text, "parts": parts, "unicode": uni}


def _display(u):
    return f"{u.first_name} {u.last_name}".strip() or u.username


def plan_send(scope="all", class_id=None, guardian_ids=None, resend=False):
    """Decide who gets a message. Returns (items, skipped) and changes nothing.
    items = [(guardian, [students])]. A student gets credentials only while she has not yet changed her password
    (unless resend), and not if a message for her is already queued or sent (unless resend)."""
    gq = Guardian.objects.all().prefetch_related("links__student__school_class")
    if scope == "selected":
        gq = gq.filter(pk__in=guardian_ids or [])
    skipped = {"already_logged_in": 0, "already_sent": 0, "no_children": 0, "inactive": 0}
    busy = {}
    for j in SmsJob.objects.filter(status__in=SmsJob.OPEN + SmsJob.DONE).only("students", "status"):
        for sid in j.students:
            busy[sid] = j.status
    items = []
    for g in gq:
        kids = []
        for link in g.links.all():
            s = link.student
            if not (s.is_active and not s.is_archived):
                skipped["inactive"] += 1
            elif scope == "class" and str(s.school_class_id) != str(class_id):
                continue
            elif not resend and not s.must_change_password:
                skipped["already_logged_in"] += 1
            elif not resend and s.pk in busy:
                skipped["already_sent"] += 1
            else:
                kids.append(s)
        if kids:
            items.append((g, kids))
        elif not g.links.exists():
            skipped["no_children"] += 1
    return items, skipped


def plan_custom(scope="all", class_id=None, guardian_ids=None):
    """Who gets a custom message (no login details): every parent with at least one active child (in the class, if one is chosen).
    Passwords and 'already sent' do not matter here. Returns (items, skipped); changes nothing."""
    gq = Guardian.objects.all().prefetch_related("links__student__school_class")
    if scope == "selected":
        gq = gq.filter(pk__in=guardian_ids or [])
    skipped = {"already_logged_in": 0, "already_sent": 0, "no_children": 0, "inactive": 0}
    items = []
    for g in gq:
        kids = [l.student for l in g.links.all() if l.student.is_active and not l.student.is_archived
                and (scope != "class" or str(l.student.school_class_id) == str(class_id))]
        if kids:
            items.append((g, kids))
        elif not g.links.exists():
            skipped["no_children"] += 1
        elif scope != "class":
            skipped["inactive"] += 1
    return items, skipped


@transaction.atomic
def queue_custom(items, user, text):
    """ONE plain message per parent. No password is created or changed, and the children are not marked as 'sent a login'."""
    for g, kids in items:
        SmsJob.objects.create(guardian=g, students=[], login=False, to_enc=g.phone_enc, body_enc=crypto.encrypt(text),
                              note=text[:160], template=text, summary="Custom message", created_by=user)
    return len(items)


@transaction.atomic
def queue_send(items, user, note="", templates=None):
    """Give each listed student a fresh random password and queue ONE message per parent."""
    templates = templates or {}
    cfg = SmsConfig.load()
    jobs, passwords = 0, {}  # a student with two parents gets ONE password, so both messages work
    for g, kids in items:
        pairs = []
        for s in kids:
            if s.pk not in passwords:
                passwords[s.pk] = new_password()
                s.set_password(passwords[s.pk])
                s.must_change_password = True
                s.save(update_fields=["password", "must_change_password"])
                Token.objects.filter(user=s).delete()  # any old login stops working
            pairs.append((_display(s), s.username, passwords[s.pk]))
        SmsJob.objects.create(
            guardian=g, students=[s.pk for s in kids], to_enc=g.phone_enc,
            body_enc=crypto.encrypt(build_body(g.language, cfg, pairs, note, templates.get(g.language, ""))), note=note,
            template=templates.get(g.language, ""),
            summary=("Login details for " + ", ".join(p[0] for p in pairs))[:200], created_by=user)
        jobs += 1
    return jobs


@transaction.atomic
def requeue(job, user=None):
    """Try a failed message again. If its text was already erased (it had been sent), make a new one with new passwords."""
    if job.status not in (SmsJob.FAILED, SmsJob.CANCELLED):
        return False
    if not job.login:  # a custom message holds no passwords: simply put the same text back in the queue
        if not job.body_enc:
            job.body_enc = crypto.encrypt(job.template or job.note)
    elif not job.body_enc or job.status == SmsJob.CANCELLED:
        kids = list(User.objects.filter(pk__in=job.students, role=User.Role.STUDENT, is_active=True))
        if not kids or not job.guardian:
            return False
        cfg, pairs = SmsConfig.load(), []
        for s in kids:
            pw = new_password()
            s.set_password(pw)
            s.must_change_password = True
            s.save(update_fields=["password", "must_change_password"])
            Token.objects.filter(user=s).delete()
            pairs.append((_display(s), s.username, pw))
        job.body_enc = crypto.encrypt(build_body(job.guardian.language, cfg, pairs, job.note, job.template))
    job.status, job.attempts, job.error, job.leased_until, job.device = SmsJob.QUEUED, 0, "", None, None
    job.save()
    return True


# ------------------------------------------------------------------ the phone
def lease_jobs(device, limit):
    now = timezone.now()
    SmsJob.objects.filter(status=SmsJob.SENDING, leased_until__lt=now, attempts__gte=MAX_ATTEMPTS).update(
        status=SmsJob.FAILED, error="The phone did not confirm sending.")
    out = []
    with transaction.atomic():
        qs = SmsJob.objects.select_for_update().filter(
            Q(status=SmsJob.QUEUED) | Q(status=SmsJob.SENDING, leased_until__lt=now, attempts__lt=MAX_ATTEMPTS))
        for j in qs.order_by("id")[:limit]:
            j.status, j.device, j.attempts = SmsJob.SENDING, device, j.attempts + 1
            j.leased_until = now + timedelta(minutes=LEASE_MINUTES)
            j.save(update_fields=["status", "device", "attempts", "leased_until"])
            out.append({"id": j.pk, "to": crypto.decrypt(j.to_enc), "body": crypto.decrypt(j.body_enc)})
    return out


def apply_report(device, job_id, status, error=""):
    j = SmsJob.objects.filter(pk=job_id).first()
    if j is None or (j.device_id and j.device_id != device.pk):
        return False
    now = timezone.now()
    if status == "sent" and j.status in (SmsJob.SENDING, SmsJob.QUEUED, SmsJob.FAILED):
        j.status, j.sent_at, j.error, j.body_enc, j.leased_until = SmsJob.SENT, now, "", "", None  # erase the passwords
    elif status == "delivered" and j.status in (SmsJob.SENT, SmsJob.SENDING, SmsJob.DELIVERED):
        j.status, j.delivered_at, j.body_enc, j.leased_until = SmsJob.DELIVERED, now, "", None
        j.sent_at = j.sent_at or now
    elif status == "failed" and j.status in (SmsJob.SENDING, SmsJob.QUEUED, SmsJob.SENT):
        j.status, j.error, j.leased_until = SmsJob.FAILED, (error or "The phone could not send it.")[:200], None
    else:
        return False
    j.save()
    return True
