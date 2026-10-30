"""Bulk student upload: parse -> validate -> report -> (admin confirms) -> create (AUTH-02..05)."""
import csv
import io
import re
import secrets
import string

from django.contrib.auth.hashers import make_password
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import transaction

from academics.models import SchoolClass
from core.models import audit
from . import temp_password
from .hashers import InitialPasswordHasher
from .models import BulkUploadBatch, User

TEMPLATE_COLUMNS = ["first name", "last name", "index number", "class or form", "stream",
                    # optional: parents are linked to the new students automatically (a parent with several children appears on several rows)
                    "parent name", "parent phone", "parent relationship", "parent 2 name", "parent 2 phone", "parent 2 relationship"]
REQUIRED_COLUMNS = TEMPLATE_COLUMNS[:4]


def _norm_header(h):
    return re.sub(r"\s+", " ", str(h or "").replace("_", " ").strip().lower())


def parse_upload(uploaded_file):
    """Return a list of dicts keyed by normalised header. Supports .csv and .xlsx."""
    name = (uploaded_file.name or "").lower()
    raw = uploaded_file.read()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
    elif name.endswith(".csv"):
        text = raw.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text)))
    else:
        raise ValueError("Please upload the template as an .xlsx or .csv file.")
    if not rows:
        raise ValueError("The file is empty.")
    headers = [_norm_header(h) for h in rows[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise ValueError("Missing column(s): " + ", ".join(missing) + ". Please use the downloadable template.")
    out = []
    for i, r in enumerate(rows[1:], start=2):  # spreadsheet row number
        if r is None or all(c in (None, "") for c in r):
            continue
        d = {h: ("" if v is None else str(v).strip()) for h, v in zip(headers, r)}
        d["_row"] = i
        out.append(d)
    return out


def validate_rows(rows):
    from guardians import services as parents  # imported here: the guardians app also imports accounts

    existing_index = set(User.objects.exclude(index_number="").values_list("index_number", flat=True))
    seen = {}
    clean, errors, warnings, with_parents = [], [], [], []
    for d in rows:
        problems = []
        first, last = d.get("first name", ""), d.get("last name", "")
        index, klass, stream = d.get("index number", ""), d.get("class or form", ""), d.get("stream", "")
        for label, val in (("first name", first), ("last name", last), ("index number", index), ("class or form", klass)):
            if not val:
                problems.append(f"missing {label}")
        if index:
            if index in seen:
                problems.append(f"duplicate index number (also on row {seen[index]})")
            elif index in existing_index:
                problems.append("index number already registered")
            seen.setdefault(index, d["_row"])
        if klass and not re.search(r"[1-4]", klass):
            problems.append("class or form must include a form number from 1 to 4")
        found, p_problems, p_warns = parents.read_parent_cells(d)
        problems += p_problems
        if problems:
            errors.append({"row": d["_row"], "problems": problems})
        else:
            row = {"first_name": first, "last_name": last, "index_number": index, "class": klass, "stream": stream}
            clean.append(row)
            if found:
                with_parents.append((d["_row"], found, row))
            if p_warns:
                warnings.append({"row": d["_row"], "problems": p_warns})
    stats = {"rows_with_parents": 0, "parents_new": 0, "parents_existing": 0, "links": 0, "parents_with_several_students": 0}
    if with_parents:
        parents.ensure_parent_key()
        bad, warn, stats = parents.settle_parents([(r, ps) for r, ps, _ in with_parents])
        for r, ps, row in with_parents:
            if r in bad:
                clean.remove(row)
                errors.append({"row": r, "problems": bad[r]})
            else:
                row["parents_enc"] = parents.encode_parents(ps)  # encrypted until the admin confirms
        for r, texts in warn.items():
            if r not in bad:
                warnings.append({"row": r, "problems": texts})
        errors.sort(key=lambda e: e["row"])
        merged = {}
        for w in warnings:
            merged.setdefault(w["row"], []).extend(w["problems"])
        warnings = [{"row": r, "problems": merged[r]} for r in sorted(merged)]
    report = {"total_rows": len(rows), "valid_rows": len(clean), "invalid_rows": len(errors), "errors": errors,
              "warnings": warnings, "parents": stats}
    return clean, report


def make_username(first_name, taken):
    """Exact first name, lowercase, no spaces; amina, amina2, amina3 ... (AUTH-04)."""
    base = re.sub(r"\s+", "", first_name.strip().lower()) or "student"
    candidate, n = base, 1
    while candidate in taken:
        n += 1
        candidate = f"{base}{n}"
    taken.add(candidate)
    return candidate


@transaction.atomic
def create_students(batch, actor):
    from guardians import services as parents

    taken = set(User.objects.values_list("username", flat=True))
    class_cache = {}
    users, wanted = [], []  # wanted[i] = the parents named on the row of users[i]
    for r in batch.rows:
        key = (r["class"], r["stream"])
        if key not in class_cache:
            level = int(re.search(r"[1-4]", r["class"]).group())
            class_cache[key], _ = SchoolClass.objects.get_or_create(
                name=r["class"], stream=r["stream"], defaults={"form_level": level}
            )
        u = User(
            username=make_username(r["first_name"], taken),
            first_name=r["first_name"],
            last_name=r["last_name"],
            index_number=r["index_number"],
            school_class=class_cache[key],
            role=User.Role.STUDENT,
        )
        # Starting password: a random-looking code (NOT the index number, which classmates know). Stored hashed; the credentials
        # sheet prints it again from the same secret, so it never sits readable in the database (AUTH-05).
        temp_password.issue(u, save=False)
        users.append(u)
        wanted.append(parents.decode_parents(r.get("parents_enc", "")))  # rows saved before this feature have none
    User.objects.bulk_create(users)
    pk_of = dict(User.objects.filter(username__in=[u.username for u in users]).values_list("username", "pk"))
    ids = list(pk_of.values())
    result = parents.link_new_students([(pk_of[u.username], ps) for u, ps in zip(users, wanted) if ps], actor)
    batch.status = BulkUploadBatch.CONFIRMED
    batch.created_user_ids = ids
    batch.rows = []  # no need to keep the data (parents' numbers included) once accounts exist
    if any(result.values()):
        batch.report = {**batch.report, "parents_result": result}
    batch.save()
    audit(actor, "students.bulk_upload", batch.pk, created=len(ids), **(result if any(result.values()) else {}))
    return ids


def validate_staff_password(password, user=None):
    """AUTH-07: teacher and admin accounts need strong passwords. Raises DjangoValidationError."""
    problems = []
    if len(password) < 10:
        problems.append("Use at least 10 characters.")
    if not (re.search(r"[a-z]", password) and re.search(r"[A-Z]", password) and re.search(r"\d", password)):
        problems.append("Use upper-case and lower-case letters and at least one number.")
    try:
        validate_password(password, user)
    except DjangoValidationError as e:
        problems += e.messages
    if problems:
        raise DjangoValidationError(problems)


def generate_password(length=12):
    """Random temporary password that always passes validate_staff_password."""
    alphabet = string.ascii_letters + string.digits
    while True:
        pw = "".join(secrets.choice(alphabet) for _ in range(length))
        if any(c.islower() for c in pw) and any(c.isupper() for c in pw) and any(c.isdigit() for c in pw):
            return pw
